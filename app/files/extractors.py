from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class FileExtractionError(RuntimeError):
    """Raised when a file cannot be converted to text."""


class UnsupportedFileTypeError(FileExtractionError):
    """Raised when the uploaded extension is outside the MVP scope."""


def validate_file_size(size_bytes: int, max_file_size_mb: int) -> None:
    max_bytes = max_file_size_mb * 1024 * 1024
    if size_bytes > max_bytes:
        raise FileExtractionError(
            f"Arquivo acima do limite de {max_file_size_mb} MB. Envie um arquivo menor."
        )


def extract_text_from_file(path: str | Path) -> str:
    file_path = Path(path)
    extension = file_path.suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError("Formato não suportado. Envie PDF, DOCX ou TXT.")
    if extension == ".txt":
        return extract_txt(file_path)
    if extension == ".docx":
        return extract_docx(file_path)
    return extract_pdf(file_path)


def extract_txt(path: str | Path) -> str:
    file_path = Path(path)
    try:
        return file_path.read_text(encoding="utf-8").strip()
    except UnicodeDecodeError:
        return file_path.read_text(encoding="latin-1").strip()


def extract_docx(path: str | Path) -> str:
    document = Document(str(path))
    paragraphs = [
        paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()
    ]
    tables: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                tables.append(" | ".join(cells))
    text = "\n".join([*paragraphs, *tables]).strip()
    if not text:
        raise FileExtractionError("Não encontrei texto extraível no DOCX.")
    return text


def extract_pdf(path: str | Path) -> str:
    reader = PdfReader(str(path))
    pages_text = [(page.extract_text() or "").strip() for page in reader.pages]
    text = "\n\n".join(item for item in pages_text if item).strip()
    if not text:
        raise FileExtractionError(
            "O PDF parece escaneado ou sem texto selecionável. Envie DOCX, TXT ou PDF com texto."
        )
    return text
